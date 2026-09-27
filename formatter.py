from docx import Document
from fpdf import FPDF
import io

def format_docx(text, doc_type):
    doc = Document()
    doc.add_heading(f"{doc_type}", 1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def format_pdf(text, doc_type):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, f"{doc_type}", ln=True)
    pdf.multi_cell(0, 10, text)
    return pdf.output(dest='S').encode('latin-1')

def format_html_preview(text):
    return f"<div style='border:1px solid #ccc; padding:20px;'><h2>LegalEase Preview</h2><p>{text}</p></div>"
